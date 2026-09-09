"""Display extraction coverage separately from independently measured accuracy."""
from __future__ import annotations

from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from cse_financial_etl.validation.acceptance import is_publishable_fact

CORE = ("PAT", "PBT", "EPS_SELECTED", "NAVPS", "OPERATING_PROFIT", "TOTAL_EQUITY",
        "TOTAL_ASSETS", "TOTAL_LIABILITIES", "TOP_LINE")

def add_accuracy_quality_sheet(wb: Workbook, *, facts: list[dict[str, Any]],
    market: list[dict[str, Any]], periods: tuple[date, ...], validation: dict[str, Any]) -> None:
    if validation.get("accuracy_basis") != "MANUAL_QA_CONTEXT":
        validation = {}
    ws = wb.create_sheet("Accuracy_Quality")
    ws.append(["Measure", "Value", "Interpretation"])
    issuers = {str(row["company_name"]) for row in market}
    dates = {p.isoformat() for p in periods}
    rows = {(r.get("issuer_name"), r.get("period_end"), r.get("metric_code")): r for r in facts
        if r.get("issuer_name") in issuers and r.get("period_end") in dates and r.get("metric_code") in CORE}
    expected = len(issuers) * len(dates) * len(CORE)
    eligible = sum(is_publishable_fact(r, release_mode="DRAFT") for r in rows.values())
    ws.append(["Expected issuer-period-metric cells", expected, "Nine financial metrics; securities sharing an issuer are counted once"])
    ws.append(["Validated draft cells", eligible, "Machine eligibility, pending institutional review where applicable"])
    ws.append(["Extraction coverage", eligible / expected if expected else None, "Coverage is not accuracy"])
    ws.cell(4, 2).number_format = "0.00%"
    ws.append(["Independent manual checks", validation.get("sample_size", 0), "MANUAL_QA fixtures only; excludes pipeline-seeded matches"])
    ws.append(["Independent manual issuers", validation.get("manual_issuer_count", 0), "Production benchmark target is independently adjudicated issuer breadth"])
    ws.append(["Manual sample accuracy", validation.get("accuracy") if validation.get("accuracy") is not None else "NOT_MEASURED", "Cannot certify the entire CSE universe"])
    ws.cell(7, 2).number_format = "0.00%"

    calibration = validation.get("certainty_calibration") or {}
    calibration_status = calibration.get("status")
    if not calibration_status:
        calibration_status = "NOT_CALIBRATED"
    ws.append(["Certainty calibration", calibration_status, "Heuristic certainty is probability-like only after independent calibration requirements are met"])
    ws.append(["Calibration sample size", calibration.get("sample_size", 0), "Independent MANUAL_QA observations with usable certainty scores"])
    ws.append(["Calibration ECE", calibration.get("expected_calibration_error") if calibration.get("expected_calibration_error") is not None else "NOT_MEASURED", "Weighted absolute gap between certainty and observed accuracy"])
    ws.append(["Calibration Brier score", calibration.get("brier_score") if calibration.get("brier_score") is not None else "NOT_MEASURED", "Mean squared error of certainty against independent correctness labels"])
    if isinstance(calibration.get("expected_calibration_error"), (int, float)):
        ws.cell(10, 2).number_format = "0.00%"
    ws.append(["Source absence accuracy", "NOT_MEASURED", "Requires adjudicated reported/absent examples"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for column, width in [("A", 38), ("B", 28), ("C", 90)]:
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
