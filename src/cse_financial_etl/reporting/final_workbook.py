"""Assemble the final human-review workbook from governed run outputs.

The normal snapshot generator is deliberately invoked only after the production/R4
outputs have been finalised. Supporting evidence is then embedded as additional
worksheets so a reviewer can inspect a run without opening a dozen separate files.
"""
from __future__ import annotations

import csv
import json
import shutil
from datetime import date
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from cse_financial_etl.reporting.excel import generate_excel

_HEADER_FILL = PatternFill("solid", fgColor="17365D")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _replace_sheet(wb: Workbook, name: str, *, index: int | None = None) -> Worksheet:
    if name in wb.sheetnames:
        del wb[name]
    return cast(Worksheet, wb.create_sheet(name, index))


def _style_table_sheet(ws: Worksheet, headers: list[str], rows: list[list[Any]]) -> None:
    ws.sheet_view.showGridLines = False
    if not headers:
        ws["A1"] = "No data available"
        return
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    for column_index, header in enumerate(headers, start=1):
        width = len(header)
        for data_row in rows[:300]:
            if column_index - 1 < len(data_row):
                value = data_row[column_index - 1]
                if value is not None:
                    width = max(width, len(str(value)))
        ws.column_dimensions[get_column_letter(column_index)].width = min(max(width + 2, 10), 42)
    for worksheet_row in ws.iter_rows(min_row=2):
        for cell in worksheet_row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def _csv_rows(path: Path) -> tuple[list[str], list[list[Any]]]:
    if not path.exists():
        return [], []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        all_rows = list(reader)
    if not all_rows:
        return [], []
    return [str(value) for value in all_rows[0]], [list(row) for row in all_rows[1:]]


def _add_csv_sheet(wb: Workbook, name: str, path: Path) -> None:
    headers, rows = _csv_rows(path)
    ws = _replace_sheet(wb, name)
    if headers:
        ws.append(headers)
        for data_row in rows:
            ws.append(data_row)
    _style_table_sheet(ws, headers, rows)


def _dict_rows(records: list[dict[str, Any]]) -> tuple[list[str], list[list[Any]]]:
    headers: list[str] = []
    seen: set[str] = set()
    for record in records:
        for key in record:
            key_text = str(key)
            if key_text not in seen:
                seen.add(key_text)
                headers.append(key_text)
    rows: list[list[Any]] = []
    for record in records:
        values: list[Any] = []
        for header in headers:
            value = record.get(header)
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
            values.append(value)
        rows.append(values)
    return headers, rows


def _add_records_sheet(wb: Workbook, name: str, records: list[dict[str, Any]]) -> None:
    headers, rows = _dict_rows(records)
    ws = _replace_sheet(wb, name)
    if headers:
        ws.append(headers)
        for data_row in rows:
            ws.append(data_row)
    _style_table_sheet(ws, headers, rows)


def _read_json(path: Path) -> Any:
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except (json.JSONDecodeError, OSError):
        return None


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    rows: list[dict[str, Any]] = []
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        if not line.strip():
            continue
        try:
            value = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(value, dict):
            rows.append(value)
    return rows


def _summary_rows(source: str, value: Any, *, prefix: str = "") -> list[list[Any]]:
    rows: list[list[Any]] = []
    if isinstance(value, dict):
        for key, child in value.items():
            nested = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(_summary_rows(source, child, prefix=nested))
    elif isinstance(value, list):
        text = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
        rows.append([source, prefix, text[:32000]])
    else:
        rows.append([source, prefix, value])
    return rows


def _add_run_summary(wb: Workbook, project_root: Path, as_of: date) -> None:
    date_text = as_of.isoformat()
    sources = {
        "Run Manifest": project_root / "outputs" / "manifests" / f"run_manifest_{date_text}.json",
        "Universe Acceptance": project_root / "outputs" / f"universe_acceptance_{date_text}.json",
        "Acceptance Recheck": project_root / "outputs" / f"universe_acceptance_recheck_{date_text}.json",
        "Row Safety": project_root / "outputs" / f"row_safety_{date_text}.json",
        "R4 Hardening": project_root / "outputs" / f"r4_hardening_{date_text}.json",
        "Quarantine Integrity": project_root / "outputs" / f"quarantine_integrity_{date_text}.json",
        "Golden Validation": project_root / "outputs" / f"golden_validation_{date_text}.json",
    }
    rows: list[list[Any]] = [
        ["Final Workbook", "assembly_stage", "POST_R4_POST_ROW_SAFETY_POST_ACCEPTANCE_RECHECK"],
        ["Final Workbook", "as_of", date_text],
        ["Final Workbook", "quarter_flow_policy", "REPORTED_THREE_MONTH_QUARTERS_ONLY"],
        ["Final Workbook", "quarter_end_price_semantic", "LAST_TRADED_ONLY"],
        ["Final Workbook", "eps_selected_policy", "DILUTED_IF_VALID_ELSE_BASIC"],
    ]
    for label, path in sources.items():
        payload = _read_json(path)
        if payload is None:
            rows.append([label, "status", "NOT_AVAILABLE"])
        else:
            rows.extend(_summary_rows(label, payload))
    ws = _replace_sheet(wb, "Run_Summary", index=1)
    headers = ["Source", "Key", "Value"]
    ws.append(headers)
    for summary_row in rows:
        ws.append(summary_row)
    _style_table_sheet(ws, headers, rows)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 90
    for worksheet_row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        worksheet_row[0].alignment = Alignment(vertical="top", wrap_text=True)


def _add_metric_definitions(wb: Workbook, path: Path) -> None:
    payload = _read_json(path)
    rows: list[list[Any]] = []
    if isinstance(payload, dict):
        for metric, definition in payload.items():
            if isinstance(definition, dict):
                for key, value in definition.items():
                    if isinstance(value, (dict, list)):
                        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
                    rows.append([metric, key, value])
            else:
                rows.append([metric, "value", definition])
    ws = _replace_sheet(wb, "Metric_Definitions")
    headers = ["Metric / Contract", "Field", "Value"]
    ws.append(headers)
    for definition_row in rows:
        ws.append(definition_row)
    _style_table_sheet(ws, headers, rows)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 90
    for worksheet_row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        worksheet_row[0].alignment = Alignment(vertical="top", wrap_text=True)


def _relabel_snapshot_headers(wb: Workbook) -> None:
    replacements = {
        "EPS": "EPS (Selected)",
        "Qtr-End Price": "Last Traded Price (Qtr End)",
        "Debt / Equity": "Liabilities / Equity",
    }
    for ws in wb.worksheets:
        if not ws.title.startswith("Snapshot_"):
            continue
        for cell in ws[4]:
            if isinstance(cell.value, str) and cell.value in replacements:
                cell.value = replacements[cell.value]
        for column in range(1, ws.max_column + 1):
            header = ws.cell(4, column).value
            if header == "Last Traded Price (Qtr End)":
                ws.column_dimensions[get_column_letter(column)].width = 18
            elif header == "Liabilities / Equity":
                ws.column_dimensions[get_column_letter(column)].width = 16
            elif header == "EPS (Selected)":
                ws.column_dimensions[get_column_letter(column)].width = 13


def _add_output_index(wb: Workbook, project_root: Path, workbook_path: Path) -> None:
    records: list[dict[str, Any]] = []
    outputs = project_root / "outputs"
    if outputs.exists():
        for path in sorted(item for item in outputs.rglob("*") if item.is_file()):
            if path == workbook_path:
                continue
            records.append(
                {
                    "relative_path": str(path.relative_to(project_root)),
                    "extension": path.suffix.lower() or "(none)",
                    "size_bytes": path.stat().st_size,
                }
            )
    _add_records_sheet(wb, "Output_Index", records)


def build_final_workbook(project_root: Path, as_of: date) -> Path:
    """Build the final consolidated workbook from the latest governed output files."""

    root = project_root.resolve()
    date_text = as_of.isoformat()
    manifest_path = root / "outputs" / "manifests" / f"run_manifest_{date_text}.json"
    manifest = _read_json(manifest_path)
    if not isinstance(manifest, dict):
        raise FileNotFoundError(f"Run manifest is required before final workbook assembly: {manifest_path}")
    run_id = str(manifest.get("run_id") or "UNKNOWN")
    raw_periods = manifest.get("target_periods") or []
    if not isinstance(raw_periods, list) or not raw_periods:
        raise ValueError("Run manifest does not contain target_periods")
    periods = tuple(date.fromisoformat(str(value)) for value in raw_periods)

    # generate_excel reads the current CSV outputs. At this stage those files have
    # already been rewritten by R4, so the visible Snapshot cannot lag hardening.
    workbook_path = generate_excel(root, as_of, periods, run_id)
    wb = load_workbook(workbook_path)
    _relabel_snapshot_headers(wb)

    _add_run_summary(wb, root, as_of)
    _add_csv_sheet(wb, "Facts_Ledger", root / "outputs" / f"normalized_facts_{date_text}.csv")
    _add_csv_sheet(wb, "Quarter_End_Prices", root / "outputs" / f"quarter_end_prices_{date_text}.csv")
    _add_csv_sheet(wb, "Review_Queue", root / "outputs" / f"review_queue_{date_text}.csv")
    _add_csv_sheet(wb, "Adjudication", root / "outputs" / f"adjudication_packet_{date_text}.csv")
    _add_metric_definitions(wb, root / "outputs" / f"metric_definitions_{date_text}.json")

    master = _read_json(root / "outputs" / f"issuer_security_master_{date_text}.json")
    if isinstance(master, dict):
        issuers = master.get("issuers") or []
        securities = master.get("securities") or []
        if isinstance(issuers, list):
            _add_records_sheet(wb, "Issuer_Master", [row for row in issuers if isinstance(row, dict)])
        if isinstance(securities, list):
            _add_records_sheet(wb, "Security_Master", [row for row in securities if isinstance(row, dict)])

    disclosures = _read_json(root / "outputs" / f"expected_disclosures_{date_text}.json")
    if isinstance(disclosures, list):
        _add_records_sheet(
            wb,
            "Expected_Disclosures",
            [row for row in disclosures if isinstance(row, dict)],
        )

    semantics = _read_jsonl(root / "outputs" / f"fact_semantics_{date_text}.jsonl")
    if semantics:
        _add_records_sheet(wb, "Fact_Semantics", semantics)

    errors = _read_json(root / "outputs" / f"pipeline_errors_{date_text}.json")
    if isinstance(errors, list):
        _add_records_sheet(wb, "Pipeline_Errors", [row for row in errors if isinstance(row, dict)])

    _add_output_index(wb, root, workbook_path)

    wb.save(workbook_path)
    compatibility_path = root / "outputs" / f"CSE_Financial_Snapshot_{date_text}.xlsx"
    if compatibility_path != workbook_path:
        shutil.copy2(workbook_path, compatibility_path)
    return workbook_path
