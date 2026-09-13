"""V2 workbook renderer. Financial cells are numeric or blank; status is elsewhere."""

from __future__ import annotations

from collections.abc import Sequence
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from cse_financial_etl.v2.contracts.facts import DerivedFact, SourceFact
from cse_financial_etl.v2.contracts.release import ReleaseContext, require_release_context
from cse_financial_etl.v2.reporting.reconciliation import (
    DisplayedWorkbookCell,
    PivotFact,
    reconcile_workbook,
)
from cse_financial_etl.v2.reporting.release_view import build_release_view

SNAPSHOT_METRICS: tuple[str, ...] = (
    "PAT",
    "PBT",
    "EPS_SELECTED",
    "NAVPS",
    "OPERATING_PROFIT",
    "TOTAL_EQUITY",
    "TOTAL_ASSETS",
    "TOTAL_LIABILITIES",
    "TOP_LINE",
    "LAST_TRADED_PRICE",
    "LIABILITIES_TO_EQUITY",
    "ROE",
    "ROA",
    "NPM",
)


def render_workbook(
    *,
    release: ReleaseContext | None,
    source_facts: Sequence[SourceFact] = (),
    derived_facts: Sequence[DerivedFact] = (),
    destination: Path,
) -> Path:
    context = require_release_context(release)
    view = build_release_view(
        release=context, source_facts=source_facts, derived_facts=derived_facts
    )
    periods = sorted({fact.period_end for fact in view.eligible})
    period: date | None = periods[-1] if periods else None
    snapshot_facts = tuple(
        fact
        for fact in view.eligible
        if fact.metric_code in SNAPSHOT_METRICS and (period is None or fact.period_end == period)
    )
    pivot = tuple(
        PivotFact(
            fact_id=fact.fact_id,
            issuer_id=fact.issuer_id,
            metric_code=fact.metric_code,
            period_end=fact.period_end,
            normalized_value=fact.normalized_value,
        )
        for fact in snapshot_facts
    )
    workbook = Workbook()
    snapshot = workbook.active
    if snapshot is None:
        snapshot = workbook.create_sheet("Snapshot")
    snapshot.title = "Snapshot"
    snapshot["A1"] = (
        f"[{context.mode.value} - not an official release]"
        if context.mode.value == "DRAFT"
        else "OFFICIAL"
    )
    snapshot["A3"] = "Issuer"
    for index, metric in enumerate(SNAPSHOT_METRICS, start=2):
        snapshot.cell(row=3, column=index, value=metric)
    issuers = sorted({fact.issuer_id for fact in view.eligible})
    displayed: list[DisplayedWorkbookCell] = []
    lookup = {(fact.issuer_id, fact.metric_code, fact.period_end): fact for fact in snapshot_facts}
    for row_index, issuer in enumerate(issuers, start=4):
        snapshot.cell(row=row_index, column=1, value=issuer)
        for col_index, metric in enumerate(SNAPSHOT_METRICS, start=2):
            fact = lookup.get((issuer, metric, period)) if period is not None else None
            value: Decimal | None = None if fact is None else fact.normalized_value
            cell = snapshot.cell(
                row=row_index, column=col_index, value=float(value) if value is not None else None
            )
            if period is not None:
                displayed.append(
                    DisplayedWorkbookCell(
                        issuer_id=issuer,
                        metric_code=metric,
                        period_end=period,
                        value=value,
                        raw_written=cell.value,
                        sheet="Snapshot",
                        cell_address=cell.coordinate,
                    )
                )

    facts_sheet = workbook.create_sheet("Financial_Facts")
    facts_sheet.append(["issuer_id", "metric_code", "period_end", "value", "entity_scope"])
    for fact in view.eligible:
        facts_sheet.append(
            [
                fact.issuer_id,
                fact.metric_code,
                fact.period_end.isoformat(),
                float(fact.normalized_value),
                fact.entity_scope.value,
            ]
        )

    missing = workbook.create_sheet("Missing_Values")
    missing.append(["reason_code"])
    for reason in view.withheld_reason_codes:
        missing.append([reason])

    review = workbook.create_sheet("Review_Summary")
    review.append(["status", "detail"])
    review.append(["REVIEW", "status strings belong on this sheet, not Snapshot"])

    quality = workbook.create_sheet("Data_Quality")
    quality.append(["eligible_facts", len(view.eligible)])

    manifest = workbook.create_sheet("Run_Manifest")
    for key, value in context.model_dump().items():
        manifest.append([key, str(value)])

    numeric_displayed = tuple(cell for cell in displayed if cell.value is not None)
    report = reconcile_workbook(
        release=context,
        eligible=snapshot_facts,
        pivot=pivot,
        displayed=numeric_displayed,
    )
    report.raise_if_failed()
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination
