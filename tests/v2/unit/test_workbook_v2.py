from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from cse_financial_etl.v2.contracts.enums import PublicationStatus, ReleaseMode, ReviewStatus, ValidationStatus
from cse_financial_etl.v2.exceptions import ReleaseContextRequiredError
from cse_financial_etl.v2.reporting.workbook import render_workbook
from tests.v2.helpers import release_context, source_fact


def test_workbook_writes_numeric_cells_and_reconciles(tmp_path: Path) -> None:
    fact = source_fact(
        validation_status=ValidationStatus.PASSED,
        publication_status=PublicationStatus.ELIGIBLE,
        normalized_value=Decimal("1234"),
    )
    path = render_workbook(
        release=release_context(),
        source_facts=(fact,),
        destination=tmp_path / "draft.xlsx",
    )
    workbook = load_workbook(path)
    assert "Snapshot" in workbook.sheetnames
    assert "Missing_Values" in workbook.sheetnames
    assert "Review_Summary" in workbook.sheetnames
    snapshot = workbook["Snapshot"]
    values = [cell.value for row in snapshot.iter_rows(min_row=4, max_row=4) for cell in row]
    assert 1234.0 in values or 1234 in values
    assert "ENTITY_NOT_RESOLVED" not in {cell.value for row in snapshot.iter_rows() for cell in row}


def test_zero_is_a_number_not_a_missing_status(tmp_path: Path) -> None:
    fact = source_fact(
        validation_status=ValidationStatus.PASSED,
        publication_status=PublicationStatus.ELIGIBLE,
        normalized_value=Decimal("0"),
        raw_value=Decimal("0"),
        period_end=date(2026, 6, 30),
    )
    path = render_workbook(
        release=release_context(), source_facts=(fact,), destination=tmp_path / "zero.xlsx"
    )
    snapshot = load_workbook(path)["Snapshot"]
    numeric = [
        cell.value
        for row in snapshot.iter_rows(min_row=4, max_row=4)
        for cell in row
        if isinstance(cell.value, (int, float))
    ]
    assert 0 in numeric or 0.0 in numeric


def test_workbook_requires_explicit_release_context(tmp_path: Path) -> None:
    with pytest.raises(ReleaseContextRequiredError):
        render_workbook(
            release=None, source_facts=(source_fact(),), destination=tmp_path / "no-ctx.xlsx"
        )


def test_withheld_facts_are_not_written_as_snapshot_status_strings(tmp_path: Path) -> None:
    fact = source_fact(
        validation_status=ValidationStatus.FAILED,
        publication_status=PublicationStatus.WITHHELD,
        reason_codes=("ENTITY_NOT_RESOLVED",),
    )
    path = render_workbook(
        release=release_context(), source_facts=(fact,), destination=tmp_path / "withheld.xlsx"
    )
    snapshot = load_workbook(path)["Snapshot"]
    written = {cell.value for row in snapshot.iter_rows() for cell in row}
    assert "ENTITY_NOT_RESOLVED" not in written
    assert "WITHHELD" not in written


def test_v2_release_path_is_authoritative_without_set_release_mode(tmp_path: Path) -> None:
    from cse_financial_etl.v2.reporting import workbook as v2_workbook

    pending = source_fact(
        validation_status=ValidationStatus.PASSED,
        publication_status=PublicationStatus.ELIGIBLE,
        review_status=ReviewStatus.REVIEW,
        normalized_value=Decimal("111"),
        raw_value=Decimal("111"),
    )
    approved = source_fact(
        fact_id="fact-approved",
        cell_id="cell-approved",
        metric_code="PBT",
        validation_status=ValidationStatus.PASSED,
        publication_status=PublicationStatus.ELIGIBLE,
        review_status=ReviewStatus.APPROVED,
        normalized_value=Decimal("222"),
        raw_value=Decimal("222"),
    )
    draft_path = render_workbook(
        release=release_context(mode=ReleaseMode.DRAFT),
        source_facts=(pending, approved),
        destination=tmp_path / "draft.xlsx",
    )
    official_path = render_workbook(
        release=release_context(mode=ReleaseMode.OFFICIAL),
        source_facts=(pending, approved),
        destination=tmp_path / "official.xlsx",
    )
    draft_values = [
        row[3]
        for row in load_workbook(draft_path)["Financial_Facts"].iter_rows(
            min_row=2, values_only=True
        )
    ]
    official_values = [
        row[3]
        for row in load_workbook(official_path)["Financial_Facts"].iter_rows(
            min_row=2, values_only=True
        )
    ]
    assert 111.0 in draft_values or 111 in draft_values
    assert 222.0 in draft_values or 222 in draft_values
    assert 222.0 in official_values or 222 in official_values
    assert 111.0 not in official_values and 111 not in official_values
    assert load_workbook(draft_path)["Snapshot"]["A1"].value != "OFFICIAL"
    assert load_workbook(official_path)["Snapshot"]["A1"].value == "OFFICIAL"
    assert "set_release_mode" not in Path(v2_workbook.__file__).read_text(encoding="utf-8")


def test_publish_module_is_authoritative_without_v1_release_mode(tmp_path: Path) -> None:
    from cse_financial_etl.v2.production import publish as publish_mod
    from cse_financial_etl.v2.production.publish import publish_production_workbook

    text = Path(publish_mod.__file__).read_text(encoding="utf-8")
    assert "set_release_mode" not in text
    assert "_release_mode" not in text
    assert "cse_financial_etl.contracts.publication" not in text
    assert "cse_financial_etl.validation.acceptance" not in text
    pending = source_fact(
        validation_status=ValidationStatus.PASSED,
        publication_status=PublicationStatus.ELIGIBLE,
        review_status=ReviewStatus.REVIEW,
        normalized_value=Decimal("111"),
        raw_value=Decimal("111"),
    )
    approved = source_fact(
        fact_id="fact-approved",
        cell_id="cell-approved",
        metric_code="PBT",
        validation_status=ValidationStatus.PASSED,
        publication_status=PublicationStatus.ELIGIBLE,
        review_status=ReviewStatus.APPROVED,
        normalized_value=Decimal("222"),
        raw_value=Decimal("222"),
    )
    draft_path = publish_production_workbook(
        release=release_context(mode=ReleaseMode.DRAFT),
        source_facts=(pending, approved),
        destination=tmp_path / "prod-draft.xlsx",
    )
    official_path = publish_production_workbook(
        release=release_context(mode=ReleaseMode.OFFICIAL),
        source_facts=(pending, approved),
        destination=tmp_path / "prod-official.xlsx",
    )
    draft = load_workbook(draft_path)
    official = load_workbook(official_path)
    assert "DRAFT" in str(draft["Snapshot"]["A1"].value)
    assert official["Snapshot"]["A1"].value == "OFFICIAL"
    draft_values = [
        row[3] for row in draft["Financial_Facts"].iter_rows(min_row=2, values_only=True)
    ]
    official_values = [
        row[3] for row in official["Financial_Facts"].iter_rows(min_row=2, values_only=True)
    ]
    assert 111.0 in draft_values or 111 in draft_values
    assert 222.0 in official_values or 222 in official_values
    assert 111.0 not in official_values and 111 not in official_values


def test_fresh_process_release_context_is_publication_authority(tmp_path: Path) -> None:
    import subprocess
    import sys
    import textwrap

    destination = tmp_path / "fresh.xlsx"
    script = tmp_path / "fresh_publish.py"
    script.write_text(
        textwrap.dedent(
            f"""
            from datetime import date
            from decimal import Decimal
            from pathlib import Path
            from cse_financial_etl.v2.contracts.enums import (
                ComparisonRole,
                EntityScope,
                PublicationStatus,
                ReleaseMode,
                ReviewStatus,
                UnitDimension,
                ValidationStatus,
            )
            from cse_financial_etl.v2.contracts.facts import SourceFact
            from cse_financial_etl.v2.contracts.provenance import SourceRef
            from cse_financial_etl.v2.contracts.release import ReleaseContext
            from cse_financial_etl.v2.production.publish import publish_production_workbook
            fact = SourceFact(
                fact_id="fact-1",
                filing_version_id="fv-1",
                statement_id="stmt-1",
                cell_id="cell-1",
                issuer_id="issuer-1",
                metric_code="PAT",
                entity_scope=EntityScope.COMPANY,
                period_end=date(2026, 6, 30),
                duration_months=3,
                comparison_role=ComparisonRole.CURRENT,
                raw_value=Decimal("333"),
                normalized_value=Decimal("333"),
                currency="LKR",
                source_scale=Decimal("1"),
                unit_dimension=UnitDimension.MONETARY,
                source_ref=SourceRef(
                    filing_id="filing-1",
                    filing_version_id="fv-1",
                    source_sha256="{'a' * 64}",
                    page_number=1,
                    bbox=(10.0, 20.0, 110.0, 36.0),
                    raw_text="333",
                    parser_name="v2.test",
                    parser_version="1.0.0",
                ),
                validation_status=ValidationStatus.PASSED,
                review_status=ReviewStatus.REVIEW,
                publication_status=PublicationStatus.ELIGIBLE,
            )
            path = publish_production_workbook(
                release=ReleaseContext(
                    generation_id="gen-1",
                    run_id="run-1",
                    mode=ReleaseMode.DRAFT,
                    code_sha="abc123",
                    policy_hash="policy-1",
                    source_snapshot_id="snap-1",
                ),
                source_facts=(fact,),
                destination=Path({str(destination)!r}),
            )
            assert path.exists()
            """
        ),
        encoding="utf-8",
    )
    completed = subprocess.run(
        [sys.executable, str(script)],
        check=False,
        cwd=Path(__file__).resolve().parents[3],
        capture_output=True,
        text=True,
    )
    assert completed.returncode == 0, completed.stderr
    assert destination.is_file()
    assert "DRAFT" in str(load_workbook(destination)["Snapshot"]["A1"].value)
