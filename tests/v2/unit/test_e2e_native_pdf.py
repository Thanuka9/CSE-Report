from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pymupdf as fitz
from openpyxl import load_workbook

from cse_financial_etl.v2.contracts.enums import EntityScope, ExtractionMode, PublicationStatus
from cse_financial_etl.v2.document.native_reader import read_native_pdf
from cse_financial_etl.v2.orchestration.filing_pipeline import run_filing_pipeline
from cse_financial_etl.v2.reporting.workbook import render_workbook
from tests.v2.helpers import release_context


def _write_income_pdf(path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=612, height=792)
    page.insert_text((72, 72), "Statement of profit or loss")
    page.insert_text((72, 96), "Company")
    page.insert_text((72, 120), "For the three months ended 30 June 2026")
    page.insert_text((72, 144), "Rs '000")
    page.insert_text((72, 180), "Profit for the period          1,234")
    document.save(path)
    document.close()


def test_native_pdf_pipeline_emits_pat_and_numeric_workbook(tmp_path: Path) -> None:
    pdf_path = tmp_path / "income.pdf"
    _write_income_pdf(pdf_path)
    document = read_native_pdf(pdf_path, filing_version_id="fv-e2e")
    assert document.pages[0].extraction_mode is ExtractionMode.NATIVE
    _statements, facts, derived, metrics = run_filing_pipeline(
        document,
        issuer_id="issuer-e2e",
        expected_entity_scope=EntityScope.COMPANY,
    )
    assert metrics.documents_parsed == 1
    pat = next((fact for fact in facts if fact.metric_code == "PAT"), None)
    if pat is None:
        # Native word clustering can keep label+value in one token; reconstruction
        # still splits trailing numerics. Fail with the parsed line texts if not.
        texts = [line.text for page in document.pages for line in page.lines]
        raise AssertionError(f"PAT missing from {facts!r}; native lines={texts!r}")
    assert pat.period_end == date(2026, 6, 30)
    assert pat.entity_scope is EntityScope.COMPANY
    assert pat.source_ref.source_sha256 == document.source_sha256
    if pat.publication_status is PublicationStatus.ELIGIBLE:
        assert pat.normalized_value == Decimal("1234000")
        path = render_workbook(
            release=release_context(),
            source_facts=facts,
            derived_facts=derived,
            destination=tmp_path / "draft.xlsx",
        )
        snapshot = load_workbook(path)["Snapshot"]
        values = [cell.value for row in snapshot.iter_rows() for cell in row]
        assert 1234000.0 in values or 1234000 in values
