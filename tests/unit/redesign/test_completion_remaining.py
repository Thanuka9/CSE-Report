"""Remaining R2 integration gates: bounded work, cache, generations and quality."""
import json
from dataclasses import replace
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from cse_financial_etl.resolution.beam_search import beam_search_concept
from cse_financial_etl.resolution.resource_budget import ResourceBudget
from cse_financial_etl.resolution.constraint_resolver import resolve_ambiguities
from cse_financial_etl.resolution.candidate_ledger import CandidateLedger
from cse_financial_etl.resolution.arbiter import arbitrate_candidates
from cse_financial_etl.storage.gold_snapshot import activate_gold_snapshot, current_gold_dir
from cse_financial_etl.ingestion.native_cache import cached_native_document
from tests.unit.redesign.test_eligibility_publisher import _entry


def test_beam_width_one_does_not_manufacture_unique_winner():
    result = beam_search_concept([_entry(), _entry(entry_id="e2", normalized_value=Decimal(200000))],
        required_entity="COMPANY", target_duration=3, budget=ResourceBudget(beam_width=1))
    assert result.status == "ABSTAIN" and result.winner is None


def test_iteration_stop_survives_final_arbitration():
    ledger = CandidateLedger([_entry(), _entry(entry_id="e2", normalized_value=Decimal(200000))])
    budget = ResourceBudget(max_iterations=1)
    resolve_ambiguities(ledger, required_entity="COMPANY", target_period_end="2026-06-30", budget=budget)
    arbitrate_candidates(ledger, required_entity="COMPANY", target_period_end="2026-06-30")
    assert budget.iterations_used == 1 and budget.stop_reason == "ITERATION_BUDGET_EXHAUSTED"
    assert not ledger.accepted()


def test_resolver_filters_wrong_period_before_scoring():
    result = beam_search_concept([_entry(period_end="2025-06-30", score=1.0)],
        required_entity="COMPANY", target_duration=3, target_period_end="2026-06-30")
    assert result.winner is None


def _generation(root, name):
    directory = root / "gold" / "snapshots" / name
    directory.mkdir(parents=True)
    for filename in ("current_financial_facts.parquet", "current_market_prices.parquet",
                     "extraction_coverage.parquet", "accuracy_certainty.parquet"):
        (directory / filename).write_bytes(name.encode())
    return directory


def test_gold_pointer_keeps_previous_generation_on_failure(tmp_path, monkeypatch):
    from cse_financial_etl.storage import stage_cache
    first = _generation(tmp_path, "first")
    activate_gold_snapshot(tmp_path, "first")
    second = _generation(tmp_path, "second")
    with monkeypatch.context() as patch:
        def crash(*args): raise OSError("interrupted switch")
        patch.setattr(stage_cache.os, "replace", crash)
        with pytest.raises(OSError): activate_gold_snapshot(tmp_path, "second")
    assert current_gold_dir(tmp_path) == first
    activate_gold_snapshot(tmp_path, "second")
    assert current_gold_dir(tmp_path) == second


def test_incomplete_generation_cannot_activate(tmp_path):
    _generation(tmp_path, "old")
    activate_gold_snapshot(tmp_path, "old")
    new = _generation(tmp_path, "new")
    (new / "current_market_prices.parquet").unlink()
    with pytest.raises(ValueError): activate_gold_snapshot(tmp_path, "new")
    assert current_gold_dir(tmp_path).name == "old"


def test_native_cache_reuses_geometry_and_invalidates_changed_source(tmp_path, monkeypatch):
    import fitz
    from cse_financial_etl.ingestion import native_pdf
    pdf = tmp_path / "test.pdf"
    def write(text):
        doc=fitz.open();page=doc.new_page();page.insert_text((50,50),text);doc.save(pdf);doc.close()
    write("Revenue 123")
    original = native_pdf.extract_native_document
    calls=[]
    def reader(path):
        calls.append(path)
        return original(path)
    monkeypatch.setattr(native_pdf, "extract_native_document", reader)
    first=cached_native_document(pdf)
    assert cached_native_document(pdf)==first and len(calls)==1
    write("Revenue 456")
    assert cached_native_document(pdf).source_sha256!=first.source_sha256 and len(calls)==2


def test_quality_sheet_counts_missing_metrics_in_denominator(tmp_path):
    from tests.unit.test_reporting import _seed_outputs, AS_OF, PERIOD, RUN_ID
    from cse_financial_etl.reporting.excel import generate_excel
    _seed_outputs(tmp_path)
    book=load_workbook(generate_excel(tmp_path,AS_OF,[PERIOD],RUN_ID))
    quality=book["Accuracy_Quality"]
    rows={r[0]:r[1] for r in quality.iter_rows(values_only=True)}
    assert rows["Expected issuer-period-metric cells"]==9
    assert rows["Manual sample accuracy"]=="NOT_MEASURED"
    assert rows["Certainty calibration"]=="NOT_CALIBRATED"
