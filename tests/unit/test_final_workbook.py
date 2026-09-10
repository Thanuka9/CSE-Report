from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from cse_financial_etl.reporting.final_workbook import build_final_workbook


def _write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def test_final_workbook_contains_review_and_lineage_sheets(tmp_path: Path) -> None:
    as_of = date(2026, 9, 10)
    date_text = as_of.isoformat()
    periods = ["2025-12-31", "2026-03-31", "2026-06-30"]

    _write(
        tmp_path / "data" / "raw" / "api" / f"market_cap_{date_text}.json",
        json.dumps(
            [
                {
                    "company_name": "TEST PLC",
                    "symbol": "TEST.N0000",
                    "price": 10.0,
                    "issued_quantity": 1000,
                    "market_capitalization": 10000.0,
                    "market_cap_percentage": 100.0,
                }
            ]
        ),
    )
    _write(
        tmp_path / "outputs" / "manifests" / f"run_manifest_{date_text}.json",
        json.dumps({"run_id": "test-run", "target_periods": periods, "issuer_count": 1}),
    )
    _write(
        tmp_path / "outputs" / f"normalized_facts_{date_text}.csv",
        "issuer_name,period_end,metric_code,normalized_value,status,validation_status,review_status\n",
    )
    _write(
        tmp_path / "outputs" / f"quarter_end_prices_{date_text}.csv",
        "symbol,period_end,value,status\n",
    )
    _write(
        tmp_path / "outputs" / f"review_queue_{date_text}.csv",
        "issuer_name,period_end,metric_code,reason,detail\n",
    )
    _write(
        tmp_path / "outputs" / f"adjudication_packet_{date_text}.csv",
        "issuer_name,period_end,metric_code,machine_value,human_value\n",
    )
    _write(
        tmp_path / "outputs" / f"metric_definitions_{date_text}.json",
        json.dumps({"MARKET_PRICE_QUARTER_END": {"publication_rule": "last traded only"}}),
    )
    _write(
        tmp_path / "outputs" / f"issuer_security_master_{date_text}.json",
        json.dumps(
            {
                "issuers": [{"issuer_id": "CSE-ISSUER-1", "legal_name": "TEST PLC"}],
                "securities": [
                    {"security_id": 1, "issuer_id": "CSE-ISSUER-1", "symbol": "TEST.N0000"}
                ],
            }
        ),
    )
    _write(
        tmp_path / "outputs" / f"expected_disclosures_{date_text}.json",
        json.dumps([{"issuer_id": "CSE-ISSUER-1", "period_end": periods[-1]}]),
    )
    _write(
        tmp_path / "outputs" / f"fact_semantics_{date_text}.jsonl",
        json.dumps({"fact_id": "fact-1", "canonical_metric_code": "PAT"}) + "\n",
    )
    _write(tmp_path / "outputs" / f"pipeline_errors_{date_text}.json", "[]")
    for name, payload in {
        "universe_acceptance": {"acceptance": "SAFE"},
        "universe_acceptance_recheck": {"acceptance": "SAFE"},
        "row_safety": {"status": "PASS", "violation_count": 0},
        "r4_hardening": {"reported_q4_only": True, "price_semantic": "LAST_TRADED_ONLY"},
        "quarantine_integrity": {"mismatches": 0},
    }.items():
        _write(tmp_path / "outputs" / f"{name}_{date_text}.json", json.dumps(payload))

    workbook_path = build_final_workbook(tmp_path, as_of)
    assert workbook_path.exists()

    workbook = load_workbook(workbook_path, read_only=True)
    expected_sheets = {
        f"Snapshot_{date_text}",
        "Run_Summary",
        "Accuracy_Quality",
        "Facts_Ledger",
        "Quarter_End_Prices",
        "Review_Queue",
        "Adjudication",
        "Metric_Definitions",
        "Issuer_Master",
        "Security_Master",
        "Expected_Disclosures",
        "Fact_Semantics",
        "Pipeline_Errors",
        "Output_Index",
    }
    assert expected_sheets.issubset(set(workbook.sheetnames))

    headers = [cell.value for cell in workbook[f"Snapshot_{date_text}"][4]]
    assert "EPS (Selected)" in headers
    assert "Last Traded Price (Qtr End)" in headers
    assert "Liabilities / Equity" in headers

    summary_rows = list(workbook["Run_Summary"].iter_rows(values_only=True))
    assert (
        "Final Workbook",
        "assembly_stage",
        "POST_R4_POST_ROW_SAFETY_POST_ACCEPTANCE_RECHECK",
    ) in summary_rows
