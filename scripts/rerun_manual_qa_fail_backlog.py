"""Rerun hybrid V2 on the Sept-28 FAIL backlog and compare to source-review values."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import time
import urllib.request
from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from cse_financial_etl.config import load_issuers
from cse_financial_etl.v2.exceptions import OcrRouteNotEnabledError
from cse_financial_etl.v2.production.engine import extract_for_production

ROOT = Path(__file__).resolve().parents[1]
DOSSIER = ROOT / "reports" / "gold_gate"
CACHE = ROOT / "data" / "tmp" / "packet_pdfs"
PER_SHARE = {"EPS_BASIC", "EPS_DILUTED", "EPS_SELECTED", "NAVPS"}
EPS_CODES = ("EPS_SELECTED", "EPS_BASIC", "EPS_DILUTED")


def _dec(value: object) -> Decimal | None:
    text = str(value or "").strip().replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _same(actual: Decimal | None, expected: Decimal | None, *, per_share: bool) -> bool:
    if actual is None or expected is None:
        return False
    if per_share:
        return abs(actual - expected) <= Decimal("0.005")
    return actual == expected


def load_manifest(path: Path) -> dict[str, dict[str, Any]]:
    wb = load_workbook(path, data_only=True)
    ws = wb["Source_Manifest"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    out: dict[str, dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, row, strict=False))
        out[str(rec["symbol"])] = rec
    return out


def download(url: str, destination: Path, expected_sha: str) -> None:
    if destination.exists():
        digest = hashlib.sha256(destination.read_bytes()).hexdigest()
        if digest == expected_sha:
            return
    destination.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0 CSE-fail-backlog-rerun/1.0", "Accept": "application/pdf,*/*"},
    )
    last_error = ""
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req, timeout=60) as response:
                payload = response.read()
            if not payload.startswith(b"%PDF"):
                raise RuntimeError("response is not a PDF")
            digest = hashlib.sha256(payload).hexdigest()
            if expected_sha and digest != expected_sha:
                raise RuntimeError(f"sha mismatch {digest} != {expected_sha}")
            destination.write_bytes(payload)
            return
        except Exception as exc:
            last_error = f"{type(exc).__name__}: {exc}"
            time.sleep(1.0 + attempt)
    raise RuntimeError(last_error or "download failed")


def selected_eps(facts: list[Any]) -> Any | None:
    by_code = {fact.metric_code: fact for fact in facts}
    for code in EPS_CODES:
        fact = by_code.get(code)
        if fact is not None and fact.normalized_value is not None:
            return fact
    return None


def match_fact(human_metric: str, facts: list[Any]) -> Any | None:
    if human_metric == "EPS_SELECTED":
        return selected_eps(facts)
    return next((fact for fact in facts if fact.metric_code == human_metric), None)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--symbols", nargs="*", default=())
    parser.add_argument(
        "--out",
        type=Path,
        default=DOSSIER / "fail_backlog_rerun.json",
    )
    args = parser.parse_args()
    wanted = {item.strip() for item in args.symbols if item.strip()}

    with (DOSSIER / "manual_qa_final_fail_queue.csv").open(newline="", encoding="utf-8") as handle:
        fails = list(csv.DictReader(handle))
    manifest = load_manifest(DOSSIER / "CSE_100_Issuer_QA_Recovery_2026-09-27.xlsx")
    issuers = load_issuers(ROOT)

    by_symbol: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in fails:
        if wanted and row["symbol"] not in wanted:
            continue
        by_symbol[row["symbol"]].append(row)

    comparisons: list[dict[str, Any]] = []
    for symbol, rows in sorted(by_symbol.items(), key=lambda item: -len(item[1])):
        info = manifest[symbol]
        pdf_path = CACHE / f"{symbol.replace('.', '_')}_{info['period_end']}.pdf"
        print(f"DOWNLOAD {symbol} {info['source_url']}", flush=True)
        download(str(info["source_url"]), pdf_path, str(info["expected_sha256"]))
        period_end = date.fromisoformat(str(info["period_end"]))
        print(f"EXTRACT {symbol}", flush=True)
        try:
            facts = extract_for_production(
                pdf_path,
                str(info["issuer_name"]),
                symbol,
                period_end,
                engine="v2",
                issuers=issuers,
            )
        except OcrRouteNotEnabledError as exc:
            for row in rows:
                comparisons.append(
                    {
                        "symbol": symbol,
                        "issuer_name": row["issuer_name"],
                        "issuer_type": row["issuer_type"],
                        "period_end": row["period_end"],
                        "metric_code": row["metric_code"],
                        "human_value": row["human_value"],
                        "human_page": row["human_source_page"],
                        "packet_machine_status": row["machine_status"],
                        "rerun_metric": None,
                        "rerun_value": None,
                        "rerun_status": "OCR_REQUIRED_NOT_AVAILABLE",
                        "match": False,
                        "residual_class": "OCR_REQUIRED",
                        "detail": str(exc),
                    }
                )
            print(f"  OCR_REQUIRED {symbol}: {exc}", flush=True)
            continue
        except Exception as exc:
            for row in rows:
                comparisons.append(
                    {
                        "symbol": symbol,
                        "issuer_name": row["issuer_name"],
                        "issuer_type": row["issuer_type"],
                        "period_end": row["period_end"],
                        "metric_code": row["metric_code"],
                        "human_value": row["human_value"],
                        "human_page": row["human_source_page"],
                        "packet_machine_status": row["machine_status"],
                        "rerun_metric": None,
                        "rerun_value": None,
                        "rerun_status": "EXTRACT_ERROR",
                        "match": False,
                        "residual_class": "EXTRACT_ERROR",
                        "detail": f"{type(exc).__name__}: {exc}",
                    }
                )
            print(f"  EXTRACT_ERROR {symbol}: {type(exc).__name__}: {exc}", flush=True)
            continue
        for row in rows:
            fact = match_fact(row["metric_code"], facts)
            human = _dec(row["human_value"])
            actual = None if fact is None else fact.normalized_value
            ok = _same(actual, human, per_share=row["metric_code"] in PER_SHARE)
            comparisons.append(
                {
                    "symbol": symbol,
                    "issuer_name": row["issuer_name"],
                    "issuer_type": row["issuer_type"],
                    "period_end": row["period_end"],
                    "metric_code": row["metric_code"],
                    "human_value": row["human_value"],
                    "human_page": row["human_source_page"],
                    "packet_machine_status": row["machine_status"],
                    "rerun_metric": None if fact is None else fact.metric_code,
                    "rerun_value": None if actual is None else str(actual),
                    "rerun_status": None if fact is None else fact.status,
                    "rerun_page": None if fact is None else fact.source_page,
                    "rerun_entity": None if fact is None else fact.entity_scope,
                    "rerun_duration": None if fact is None else fact.duration_months,
                    "rerun_method": None if fact is None else fact.extraction_method,
                    "match": ok,
                }
            )
            mark = "PASS" if ok else "FAIL"
            print(
                f"  {mark} {row['metric_code']} human={row['human_value']} "
                f"rerun={None if actual is None else actual}",
                flush=True,
            )

    matched = sum(1 for row in comparisons if row["match"])
    summary = {
        "compared": len(comparisons),
        "matched": matched,
        "still_fail": len(comparisons) - matched,
        "issuers": len(by_symbol),
        "rows": comparisons,
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")
    print(json.dumps({"compared": summary["compared"], "matched": matched, "still_fail": summary["still_fail"]}))
    print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
